"""Employers router — dashboard overview, settings, payroll deduction management.

All endpoints require employer_admin or manager JWT role.
"""

from uuid import UUID

import structlog
from fastapi import APIRouter, Depends, HTTPException, Request
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from src.core.database import get_db
from src.models import EWATransaction, Employee, Employer, PayrollUpload
from src.schemas import DeductionReport, PayrollUploadResponse
from src.services.ewa.deduction import PayrollDeductionEngine

logger = structlog.get_logger()
router = APIRouter(prefix="/api/v1/employers", tags=["employers"])


@router.get("/dashboard")
async def dashboard_overview(
    request: Request,
    db: AsyncSession = Depends(get_db),
) -> dict:
    """Employer dashboard: withdrawals today, active employees, recovery rate."""
    employer_id = _get_employer_id(request)

    from datetime import datetime, timezone

    today = datetime.now(timezone.utc).date()
    pay_period = f"{today.year}-{today.month:02d}"

    # Active employees count
    emp_result = await db.execute(
        select(func.count()).where(
            Employee.employer_id == employer_id,
            Employee.is_active.is_(True),
        )
    )
    active_employees = emp_result.scalar() or 0

    # Today's withdrawals
    today_result = await db.execute(
        select(
            func.count(EWATransaction.id),
            func.coalesce(func.sum(EWATransaction.amount_disbursed), 0),
        ).where(
            EWATransaction.employer_id == employer_id,
            EWATransaction.status == "completed",
            func.date(EWATransaction.created_at) == today,
        )
    )
    today_row = today_result.one()
    withdrawals_today = today_row[0]
    amount_today = float(today_row[1])

    # Period totals
    period_result = await db.execute(
        select(
            func.count(EWATransaction.id),
            func.coalesce(func.sum(EWATransaction.amount_requested), 0),
            func.coalesce(func.sum(EWATransaction.fee), 0),
        ).where(
            EWATransaction.employer_id == employer_id,
            EWATransaction.pay_period == pay_period,
            EWATransaction.status == "completed",
        )
    )
    period_row = period_result.one()

    # Recovery rate (deducted / total completed this period)
    deducted_result = await db.execute(
        select(func.count()).where(
            EWATransaction.employer_id == employer_id,
            EWATransaction.pay_period == pay_period,
            EWATransaction.status == "completed",
            EWATransaction.deducted_at.isnot(None),
        )
    )
    deducted_count = deducted_result.scalar() or 0
    total_completed = period_row[0]
    recovery_rate = (deducted_count / total_completed * 100) if total_completed > 0 else 0

    return {
        "active_employees": active_employees,
        "withdrawals_today": withdrawals_today,
        "amount_today": amount_today,
        "period": {
            "pay_period": pay_period,
            "total_transactions": total_completed,
            "total_amount": float(period_row[1]),
            "total_fees": float(period_row[2]),
        },
        "recovery_rate": round(recovery_rate, 1),
    }


@router.get("/settings")
async def get_settings(
    request: Request,
    db: AsyncSession = Depends(get_db),
) -> dict:
    """Get employer settings."""
    employer_id = _get_employer_id(request)

    result = await db.execute(
        select(Employer).where(Employer.id == employer_id)
    )
    employer = result.scalar_one_or_none()
    if employer is None:
        raise HTTPException(status_code=404, detail="Employer not found")

    return {
        "company_name": employer.company_name,
        "ghana_tin": employer.ghana_tin,
        "payroll_cycle": employer.payroll_cycle,
        "payday": employer.payday,
        "ewa_enabled": employer.ewa_enabled,
        "plan": employer.plan,
        "totp_verified": employer.totp_verified,
    }


@router.patch("/settings")
async def update_settings(
    request: Request,
    ewa_enabled: bool | None = None,
    payday: int | None = None,
    db: AsyncSession = Depends(get_db),
) -> dict:
    """Update employer settings (EWA toggle, payday)."""
    employer_id = _get_employer_id(request)

    result = await db.execute(
        select(Employer).where(Employer.id == employer_id)
    )
    employer = result.scalar_one_or_none()
    if employer is None:
        raise HTTPException(status_code=404, detail="Employer not found")

    if ewa_enabled is not None:
        employer.ewa_enabled = ewa_enabled
    if payday is not None:
        if not (1 <= payday <= 31):
            raise HTTPException(status_code=400, detail="Payday must be 1-31")
        employer.payday = payday

    await db.flush()

    return {"message": "Settings updated", "ewa_enabled": employer.ewa_enabled, "payday": employer.payday}


# ── Payroll Deduction endpoints ──

@router.get("/deductions/report", response_model=DeductionReport)
async def get_deduction_report(
    request: Request,
    pay_period: str,
    db: AsyncSession = Depends(get_db),
) -> DeductionReport:
    """Generate deduction report for a pay period (preview, read-only)."""
    employer_id = _get_employer_id(request)

    engine = PayrollDeductionEngine(db)
    return await engine.generate_report(employer_id, pay_period)


@router.post("/deductions/apply")
async def apply_deductions(
    request: Request,
    pay_period: str,
    db: AsyncSession = Depends(get_db),
) -> dict:
    """Apply deductions: mark completed EWA transactions as deducted.

    Idempotent: running twice returns 0 on second run.
    """
    employer_id = _get_employer_id(request)

    # Create payroll upload record
    upload = PayrollUpload(
        employer_id=employer_id,
        uploaded_by=_get_user_id(request),
        filename=f"deductions_{pay_period}.csv",
        pay_period=pay_period,
    )
    db.add(upload)
    await db.flush()

    engine = PayrollDeductionEngine(db)
    rows_deducted = await engine.apply_deductions(
        employer_id=employer_id,
        pay_period=pay_period,
        payroll_upload_id=upload.id,
    )

    return {
        "pay_period": pay_period,
        "rows_deducted": rows_deducted,
        "payroll_upload_id": str(upload.id),
    }


@router.get("/payroll-uploads", response_model=list[PayrollUploadResponse])
async def list_payroll_uploads(
    request: Request,
    limit: int = 20,
    db: AsyncSession = Depends(get_db),
) -> list[PayrollUploadResponse]:
    """List payroll upload records."""
    employer_id = _get_employer_id(request)

    result = await db.execute(
        select(PayrollUpload)
        .where(PayrollUpload.employer_id == employer_id)
        .order_by(PayrollUpload.created_at.desc())
        .limit(limit)
    )
    uploads = result.scalars().all()
    return [PayrollUploadResponse.model_validate(u) for u in uploads]


@router.get("/deductions/export")
async def export_deduction_report_excel(
    request: Request,
    pay_period: str,
    db: AsyncSession = Depends(get_db),
):
    """Download deduction report as Excel (.xlsx) file.

    Generates an openpyxl workbook with per-employee deduction breakdown.
    """
    from fastapi.responses import StreamingResponse
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    employer_id = _get_employer_id(request)
    engine = PayrollDeductionEngine(db)
    report = await engine.generate_report(employer_id, pay_period)

    # Build workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Deductions {pay_period}"

    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0A5C34", end_color="0A5C34", fill_type="solid")
    money_fmt = '#,##0.00'
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # Title
    ws.merge_cells("A1:F1")
    ws["A1"] = f"WageNow — Payroll Deduction Report: {pay_period}"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)
    ws["A2"] = f"Generated: {__import__('datetime').datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(name="Calibri", size=10, color="666666")

    # Headers
    headers = ["Employee Name", "Employee ID", "Total Withdrawn (GHS)", "Total Fees (GHS)", "Net Deduction (GHS)", "Transactions"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Data rows
    for row_idx, item in enumerate(report.items, 5):
        ws.cell(row=row_idx, column=1, value=item.employee_name).border = thin_border
        ws.cell(row=row_idx, column=2, value=str(item.employee_id)[:8]).border = thin_border
        c3 = ws.cell(row=row_idx, column=3, value=float(item.total_withdrawn))
        c3.number_format = money_fmt
        c3.border = thin_border
        c4 = ws.cell(row=row_idx, column=4, value=float(item.total_fees))
        c4.number_format = money_fmt
        c4.border = thin_border
        c5 = ws.cell(row=row_idx, column=5, value=float(item.net_deduction))
        c5.number_format = money_fmt
        c5.border = thin_border
        c5.font = Font(bold=True)
        ws.cell(row=row_idx, column=6, value=item.transaction_count).border = thin_border

    # Totals row
    total_row = len(report.items) + 5
    ws.cell(row=total_row, column=1, value="GRAND TOTAL").font = Font(bold=True)
    ct = ws.cell(row=total_row, column=5, value=float(report.grand_total))
    ct.number_format = money_fmt
    ct.font = Font(bold=True, size=12)
    ws.cell(row=total_row, column=4, value=float(report.total_fees)).number_format = money_fmt
    ws.cell(row=total_row, column=6, value=report.transaction_count)

    # Column widths
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 14

    # Write to buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"wagenow_deductions_{pay_period}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/analytics")
async def get_analytics(
    request: Request,
    months: int = 6,
    db: AsyncSession = Depends(get_db),
) -> dict:
    """Withdrawal analytics: monthly trends, department breakdown, cost savings.

    Returns data for the analytics dashboard.
    """
    from datetime import datetime, timezone
    from dateutil.relativedelta import relativedelta

    employer_id = _get_employer_id(request)
    now = datetime.now(timezone.utc)

    # Monthly trend (last N months)
    monthly_data = []
    for i in range(months - 1, -1, -1):
        d = now - relativedelta(months=i)
        period = f"{d.year}-{d.month:02d}"
        result = await db.execute(
            select(
                func.count(EWATransaction.id),
                func.coalesce(func.sum(EWATransaction.amount_disbursed), 0),
                func.coalesce(func.sum(EWATransaction.fee), 0),
            ).where(
                EWATransaction.employer_id == employer_id,
                EWATransaction.pay_period == period,
                EWATransaction.status == "completed",
            )
        )
        row = result.one()
        monthly_data.append({
            "period": period,
            "month": d.strftime("%b"),
            "transactions": row[0],
            "amount": float(row[1]),
            "fees": float(row[2]),
        })

    # Department breakdown
    dept_result = await db.execute(
        select(
            Employee.department,
            func.count(func.distinct(EWATransaction.employee_id)).label("employees"),
            func.coalesce(func.sum(EWATransaction.amount_disbursed), 0).label("amount"),
        )
        .join(Employee, EWATransaction.employee_id == Employee.id)
        .where(
            EWATransaction.employer_id == employer_id,
            EWATransaction.status == "completed",
        )
        .group_by(Employee.department)
        .order_by(func.sum(EWATransaction.amount_disbursed).desc())
    )
    dept_rows = dept_result.all()
    total_amount = sum(float(r.amount) for r in dept_rows) or 1

    departments = [
        {
            "department": r.department or "Other",
            "employees": r.employees,
            "amount": float(r.amount),
            "percentage": round(float(r.amount) / total_amount * 100, 1),
        }
        for r in dept_rows
    ]

    # Summary stats
    total_disbursed = sum(m["amount"] for m in monthly_data)
    total_txns = sum(m["transactions"] for m in monthly_data)
    avg_withdrawal = total_disbursed / total_txns if total_txns > 0 else 0
    # Estimated savings: avg Fido loan fee (GHS 60) vs WageNow fee (GHS 3) per withdrawal
    estimated_savings = total_txns * (60 - 3)

    return {
        "monthly_trend": monthly_data,
        "departments": departments,
        "summary": {
            "total_disbursed": round(total_disbursed, 2),
            "total_transactions": total_txns,
            "avg_withdrawal": round(avg_withdrawal, 2),
            "estimated_savings": round(estimated_savings, 2),
        },
    }


def _get_employer_id(request: Request) -> UUID:
    employer_id = getattr(request.state, "employer_id", None)
    if employer_id is None:
        raise HTTPException(status_code=401, detail="Authentication required")
    return UUID(employer_id) if isinstance(employer_id, str) else employer_id


def _get_user_id(request: Request) -> UUID:
    sub = getattr(request.state, "sub", None)
    if sub is None:
        raise HTTPException(status_code=401, detail="Authentication required")
    return UUID(sub)
